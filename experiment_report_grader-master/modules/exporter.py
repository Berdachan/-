"""
数据导出模块（优化版）
负责：Excel导出、PDF导出、中文字体处理
"""

from pathlib import Path
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


class ResultExporter:
    """结果导出器"""

    def __init__(self):
        self._register_chinese_fonts()

    def _register_chinese_fonts(self):
        """注册中文字体，解决PDF乱码问题"""
        self.chinese_font = "SimSun"
        self.chinese_font_bold = "SimSun-Bold"

        font_paths = [
            ("SimSun", r"C:\Windows\Fonts\simsun.ttc"),
            ("SimSun-Bold", r"C:\Windows\Fonts\simsun.ttc"),
            ("MicrosoftYaHei", r"C:\Windows\Fonts\msyh.ttc"),
            ("MicrosoftYaHei-Bold", r"C:\Windows\Fonts\msyhbd.ttc"),
        ]

        for font_name, font_path in font_paths:
            try:
                pdfmetrics.registerFont(TTFont(font_name, font_path))
            except Exception:
                pass

        try:
            pdfmetrics.getFont("SimSun")
        except Exception:
            try:
                pdfmetrics.getFont("MicrosoftYaHei")
                self.chinese_font = "MicrosoftYaHei"
                self.chinese_font_bold = "MicrosoftYaHei-Bold"
            except Exception:
                self.chinese_font = "Helvetica"
                self.chinese_font_bold = "Helvetica-Bold"

    def export_excel(self, results: list[dict], filenames: list[str], output_path: Path):
        """导出Excel"""
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "批阅结果"

        title_font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
        title_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        title_alignment = Alignment(horizontal="center", vertical="center")

        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        cell_font = Font(name="微软雅黑", size=10)
        cell_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell_alignment_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )

        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = "实验报告批阅结果"
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = title_alignment
        ws.row_dimensions[1].height = 35

        ws.merge_cells("A2:H2")
        date_cell = ws["A2"]
        date_cell.value = f"生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        date_cell.font = Font(name="微软雅黑", size=10, italic=True, color="666666")
        date_cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[2].height = 25

        headers = ["文件名", "总分", "实验目的", "实验步骤", "结果分析", "结论", "总体评语", "总体建议"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        ws.row_dimensions[4].height = 30

        for i, (result, filename) in enumerate(zip(results, filenames), 1):
            row = i + 4

            cell = ws.cell(row=row, column=1, value=filename)
            cell.font = cell_font
            cell.alignment = cell_alignment_center
            cell.border = thin_border

            cell = ws.cell(row=row, column=2, value=f"{result['total_score']}/{result['total_max']}")
            cell.font = Font(name="微软雅黑", size=11, bold=True, color="C00000")
            cell.alignment = cell_alignment_center
            cell.border = thin_border

            dims = result["dimensions"]
            dim_cols = ["实验目的理解", "实验步骤完整性", "结果分析与数据处理", "结论与总结"]
            for col_idx, dim_name in enumerate(dim_cols, 3):
                dim = dims.get(dim_name, {})
                score = dim.get("score", "-")
                max_score = dim.get("max_score", "-")
                comment = dim.get("comment", "无评语")
                suggestions = dim.get("suggestions", "无建议")

                cell_value = f"得分: {score}/{max_score}\n\n评语: {comment}\n\n建议: {suggestions}"
                cell = ws.cell(row=row, column=col_idx, value=cell_value)
                cell.font = cell_font
                cell.alignment = cell_alignment
                cell.border = thin_border

            cell = ws.cell(row=row, column=7, value=result["overall_comment"])
            cell.font = cell_font
            cell.alignment = cell_alignment
            cell.border = thin_border

            cell = ws.cell(row=row, column=8, value=result["overall_suggestions"])
            cell.font = cell_font
            cell.alignment = cell_alignment
            cell.border = thin_border

            ws.row_dimensions[row].height = 120

        col_widths = {"A": 25, "B": 12, "C": 25, "D": 25, "E": 25, "F": 25, "G": 35, "H": 35}
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = "A5"
        wb.save(output_path)

    def export_pdf(self, results: list[dict], filenames: list[str], output_path: Path):
        """导出PDF"""
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            rightMargin=1.5*cm,
            leftMargin=1.5*cm,
            topMargin=1.5*cm,
            bottomMargin=1.5*cm
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Heading1"],
            fontName=self.chinese_font_bold,
            fontSize=18,
            alignment=1,
            spaceAfter=15,
            textColor=colors.HexColor("#1a5276")
        )

        heading2_style = ParagraphStyle(
            "CustomHeading2",
            parent=styles["Heading2"],
            fontName=self.chinese_font_bold,
            fontSize=12,
            spaceAfter=8,
            textColor=colors.HexColor("#2874a6")
        )

        normal_style = ParagraphStyle(
            "CustomNormal",
            parent=styles["Normal"],
            fontName=self.chinese_font,
            fontSize=9,
            leading=14,
            spaceAfter=6
        )

        bold_style = ParagraphStyle(
            "CustomBold",
            parent=styles["Normal"],
            fontName=self.chinese_font_bold,
            fontSize=9,
            leading=14,
            spaceAfter=6,
            textColor=colors.HexColor("#1a5276")
        )

        score_style = ParagraphStyle(
            "ScoreStyle",
            parent=styles["Normal"],
            fontName=self.chinese_font_bold,
            fontSize=14,
            leading=18,
            spaceAfter=10,
            textColor=colors.HexColor("#c0392b")
        )

        cell_header_style = ParagraphStyle(
            "CellHeader",
            fontName=self.chinese_font_bold,
            fontSize=9,
            leading=12,
            alignment=1,
            textColor=colors.whitesmoke,
        )

        cell_style = ParagraphStyle(
            "CellStyle",
            fontName=self.chinese_font,
            fontSize=8,
            leading=11,
            alignment=0,
            wordWrap='CJK',
        )

        cell_style_center = ParagraphStyle(
            "CellStyleCenter",
            fontName=self.chinese_font,
            fontSize=8,
            leading=11,
            alignment=1,
            wordWrap='CJK',
        )

        story = []
        story.append(Paragraph("实验报告批阅结果", title_style))
        story.append(Spacer(1, 0.2*cm))
        story.append(Paragraph(
            f"<font name='{self.chinese_font}' size='8' color='#666666'>生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</font>",
            normal_style
        ))
        story.append(Spacer(1, 0.3*cm))
        story.append(Table([[""]], colWidths=[18*cm], style=TableStyle([
            ("LINEBELOW", (0, 0), (-1, 0), 1.5, colors.HexColor("#1a5276")),
        ])))
        story.append(Spacer(1, 0.3*cm))

        for i, (result, filename) in enumerate(zip(results, filenames), 1):
            story.append(Paragraph(f"报告: {filename}", heading2_style))
            story.append(Spacer(1, 0.15*cm))
            story.append(Paragraph(
                f"总分: {result['total_score']}/{result['total_max']} 分  (得分率: {result['percentage']}%)",
                score_style
            ))
            story.append(Spacer(1, 0.15*cm))

            data = [[
                Paragraph("维度", cell_header_style),
                Paragraph("得分", cell_header_style),
                Paragraph("评语", cell_header_style),
                Paragraph("改进建议", cell_header_style)
            ]]

            for name, dim in result["dimensions"].items():
                data.append([
                    Paragraph(name, cell_style_center),
                    Paragraph(f"{dim['score']}/{dim['max_score']}", cell_style_center),
                    Paragraph(dim["comment"], cell_style),
                    Paragraph(dim["suggestions"], cell_style)
                ])

            table = Table(data, colWidths=[3.2*cm, 1.8*cm, 6.5*cm, 6.5*cm])
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a5276")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("FONTNAME", (0, 0), (-1, 0), self.chinese_font_bold),
                ("FONTNAME", (0, 1), (-1, -1), self.chinese_font),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("FONTSIZE", (0, 1), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 10),
                ("TOPPADDING", (0, 1), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
                ("LEFTPADDING", (0, 1), (-1, -1), 6),
                ("RIGHTPADDING", (0, 1), (-1, -1), 6),
                ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f8f9fa")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#bdc3c7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.3*cm))

            story.append(Paragraph("总体评语:", bold_style))
            story.append(Paragraph(result["overall_comment"], normal_style))
            story.append(Spacer(1, 0.15*cm))

            story.append(Paragraph("总体建议:", bold_style))
            story.append(Paragraph(result["overall_suggestions"], normal_style))
            story.append(Spacer(1, 0.15*cm))

            if result.get("strengths"):
                story.append(Paragraph("优点:", bold_style))
                for s in result["strengths"]:
                    story.append(Paragraph(f"• {s}", normal_style))
                story.append(Spacer(1, 0.15*cm))

            if result.get("weaknesses"):
                story.append(Paragraph("不足:", bold_style))
                for w in result["weaknesses"]:
                    story.append(Paragraph(f"• {w}", normal_style))
                story.append(Spacer(1, 0.15*cm))

            story.append(Table([[""]], colWidths=[18*cm], style=TableStyle([
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#bdc3c7")),
            ])))
            story.append(Spacer(1, 0.3*cm))

            if i < len(results):
                story.append(PageBreak())

        doc.build(story)

    def export_ppt(self, results: list[dict], filenames: list[str], output_path: Path):
        """导出PPT"""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
            
            prs = Presentation()
            
            # 设置幻灯片尺寸为16:9
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            
            # 标题页
            title_slide_layout = prs.slide_layouts[6]  # 空白布局
            slide = prs.slides.add_slide(title_slide_layout)
            
            # 添加标题
            title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
            tf = title_box.text_frame
            p = tf.paragraphs[0]
            p.text = "实验报告批阅结果"
            p.font.size = Pt(44)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0x1a, 0x52, 0x76)
            p.alignment = PP_ALIGN.CENTER
            
            # 添加日期
            date_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(0.5))
            tf = date_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            p.font.size = Pt(16)
            p.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            p.alignment = PP_ALIGN.CENTER
            
            # 内容页
            for i, (result, filename) in enumerate(zip(results, filenames), 1):
                slide_layout = prs.slide_layouts[6]
                slide = prs.slides.add_slide(slide_layout)
                
                # 文件名标题
                title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
                tf = title_shape.text_frame
                p = tf.paragraphs[0]
                p.text = f"{i}. {filename}"
                p.font.size = Pt(24)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0x28, 0x74, 0xa6)
                
                # 总分
                score_shape = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(12.333), Inches(0.6))
                tf = score_shape.text_frame
                p = tf.paragraphs[0]
                p.text = f"总分: {result['total_score']}/{result['total_max']} ({result['percentage']}%)"
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = RGBColor(0xc0, 0x39, 0x2b)
                
                # 各维度
                y_pos = 2.0
                for name, dim in result["dimensions"].items():
                    dim_shape = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(6), Inches(1.2))
                    tf = dim_shape.text_frame
                    tf.word_wrap = True
                    
                    p = tf.paragraphs[0]
                    p.text = f"{name}: {dim['score']}/{dim['max_score']}"
                    p.font.size = Pt(14)
                    p.font.bold = True
                    
                    p = tf.add_paragraph()
                    p.text = f"评语: {dim['comment']}"
                    p.font.size = Pt(12)
                    
                    y_pos += 1.3
                
                # 总体评语
                if y_pos < 6:
                    comment_shape = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(12.333), Inches(1.5))
                    tf = comment_shape.text_frame
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    p.text = "总体评语:"
                    p.font.size = Pt(14)
                    p.font.bold = True
                    p = tf.add_paragraph()
                    p.text = result["overall_comment"][:200] + "..." if len(result["overall_comment"]) > 200 else result["overall_comment"]
                    p.font.size = Pt(12)
            
            prs.save(str(output_path))
            
        except ImportError:
            raise ImportError("python-pptx 未安装，请运行: pip install python-pptx")
        
    def export_docx(self, results: list[dict], filenames: list[str], output_path: Path):
        """导出Word文档（优化版，解决中文格式混乱）"""
        try:
            import importlib.util

            if importlib.util.find_spec("docx") is None:
                raise ImportError

            from docx import Document
            from docx.shared import Pt, RGBColor, Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
            
            doc = Document()

            # 设置页面边距（减小边距给表格更多空间）
            sections = doc.sections[0]
            sections.left_margin = Cm(1.5)
            sections.right_margin = Cm(1.5)
            sections.top_margin = Cm(2)
            sections.bottom_margin = Cm(2)
            
            # ========== 设置全局中文字体 ==========
            # 设置默认字体
            style = doc.styles['Normal']
            font = style.font
            font.name = '微软雅黑'
            font.size = Pt(10.5)
            # 设置中文字体
            style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            
            # 设置段落默认样式
            paragraph_format = style.paragraph_format
            paragraph_format.line_spacing = Pt(18)  # 行距
            paragraph_format.space_after = Pt(6)    # 段后间距
            
            # ========== 标题 ==========
            title = doc.add_heading("实验报告批阅结果", level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # 设置标题字体
            for run in title.runs:
                run.font.name = '微软雅黑'
                run.font.size = Pt(22)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1a, 0x52, 0x76)
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            
            # 日期
            date_para = doc.add_paragraph()
            date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = date_para.add_run(f"生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            run.font.name = '微软雅黑'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
            
            doc.add_paragraph()  # 空行
            
            # ========== 内容 ==========
            for i, (result, filename) in enumerate(zip(results, filenames), 1):
                # 文件名标题
                heading = doc.add_heading(f"{i}. {filename}", level=1)
                for run in heading.runs:
                    run.font.name = '微软雅黑'
                    run.font.size = Pt(16)
                    run.font.color.rgb = RGBColor(0x28, 0x74, 0xa6)
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                
                # 总分
                score_para = doc.add_paragraph()
                run = score_para.add_run(f"总分: {result['total_score']}/{result['total_max']} ({result['percentage']}%)")
                run.font.name = '微软雅黑'
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xc0, 0x39, 0x2b)
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                
                # 各维度表格
                doc.add_paragraph()
                table = doc.add_table(rows=1, cols=4)
                table.style = 'Table Grid'

                # 设置表格自动调整
                table.autofit = True  # 改为 True，让表格自动适应内容

                # 表头
                hdr_cells = table.rows[0].cells
                headers = ["维度", "得分", "评语", "建议"]
                for j, header in enumerate(headers):
                    hdr_cells[j].text = header
                    for paragraph in hdr_cells[j].paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in paragraph.runs:
                            run.font.name = '微软雅黑'
                            run.font.size = Pt(11)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

                # 数据行
                for name, dim in result["dimensions"].items():
                    row_cells = table.add_row().cells
                    row_cells[0].text = name
                    row_cells[1].text = f"{dim['score']}/{dim['max_score']}"
                    row_cells[2].text = dim["comment"]
                    row_cells[3].text = dim["suggestions"]
                    
                    # 设置单元格垂直居中
                    for cell in row_cells:
                        cell.vertical_alignment = WD_ALIGN_PARAGRAPH.CENTER  # 垂直居中
                        for paragraph in cell.paragraphs:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT  # 水平左对齐
                            for run in paragraph.runs:
                                run.font.name = '微软雅黑'
                                run.font.size = Pt(9)  # 稍微小一点
                                run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                                
                    doc.add_paragraph()  # 空行
                
                # 总体评语
                comment_heading = doc.add_heading("总体评语", level=2)
                for run in comment_heading.runs:
                    run.font.name = '微软雅黑'
                    run.font.size = Pt(13)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x28, 0x74, 0xa6)
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                
                comment_para = doc.add_paragraph(result["overall_comment"])
                for run in comment_para.runs:
                    run.font.name = '微软雅黑'
                    run.font.size = Pt(10.5)
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                
                # 总体建议
                suggestions_heading = doc.add_heading("总体建议", level=2)
                for run in suggestions_heading.runs:
                    run.font.name = '微软雅黑'
                    run.font.size = Pt(13)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x28, 0x74, 0xa6)
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                
                suggestions_para = doc.add_paragraph(result["overall_suggestions"])
                for run in suggestions_para.runs:
                    run.font.name = '微软雅黑'
                    run.font.size = Pt(10.5)
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                
                # 优点
                if result.get("strengths"):
                    strengths_heading = doc.add_heading("优点", level=2)
                    for run in strengths_heading.runs:
                        run.font.name = '微软雅黑'
                        run.font.size = Pt(13)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0x22, 0x8B, 0x22)
                        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                    
                    for s in result["strengths"]:
                        p = doc.add_paragraph(s, style='List Bullet')
                        for run in p.runs:
                            run.font.name = '微软雅黑'
                            run.font.size = Pt(10.5)
                            run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                
                # 不足
                if result.get("weaknesses"):
                    weaknesses_heading = doc.add_heading("不足", level=2)
                    for run in weaknesses_heading.runs:
                        run.font.name = '微软雅黑'
                        run.font.size = Pt(13)
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xDC, 0x14, 0x3C)
                        run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                    
                    for w in result["weaknesses"]:
                        p = doc.add_paragraph(w, style='List Bullet')
                        for run in p.runs:
                            run.font.name = '微软雅黑'
                            run.font.size = Pt(10.5)
                            run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                
                # 分页
                if i < len(results):
                    doc.add_page_break()
            
            doc.save(str(output_path))
            
        except ImportError:
            raise ImportError("python-docx 未安装，请运行: pip install python-docx")