"""
历史记录管理模块
负责：查看历史记录、导出历史记录Excel、重新批阅覆盖记录
"""

import logging
from pathlib import Path
from typing import List, Optional, Dict, Any
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from modules.database import DatabaseManager, GradingRecord
from modules.exporter import ResultExporter

logger = logging.getLogger(__name__)


class HistoryManager:
    """历史记录管理器"""

    def __init__(self, db: DatabaseManager):
        self.db = db
        self.exporter = ResultExporter()

    def get_user_history(self, user_id: int, page: int = 1, page_size: int = 20) -> Dict[str, Any]:
        """获取用户历史记录（分页）"""
        offset = (page - 1) * page_size
        records = self.db.get_user_records(user_id, limit=page_size, offset=offset)
        stats = self.db.get_user_statistics(user_id)

        return {
            'records': records,
            'statistics': stats,
            'page': page,
            'page_size': page_size,
            'has_more': len(records) == page_size
        }

    def format_record_display(self, record: GradingRecord) -> str:
        """格式化单条记录为显示文本"""
        lines = [
            "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━",
            f"  ID: {record.display_id} | 文件: {record.filename}",
            f"  评分: {record.total_score}/{record.total_max} ({record.percentage}%)",
            f"  风格: {record.style_name}",
            f"  时间: {record.created_at.strftime('%Y-%m-%d %H:%M:%S')}",
            f"  质量: {record.quality_score:.1f}" if record.quality_score else "  质量: 未检测",
        ]

        if record.dimensions:
            lines.append("  ┌──────────────────────────────────────")
            for name, dim in record.dimensions.items():
                score = dim.get('score', 0)
                max_s = dim.get('max_score', 0)
                lines.append(f"  │ {name}: {score}/{max_s}")
            lines.append("  └──────────────────────────────────────")

        if record.overall_comment:
            comment = record.overall_comment[:100] + "..." if len(record.overall_comment) > 100 else record.overall_comment
            lines.append(f"  评语: {comment}")

        lines.append("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
        return "\n".join(lines)

    def format_history_summary(self, records: List[GradingRecord]) -> str:
        """格式化历史记录列表摘要"""
        if not records:
            return "暂无历史记录"

        lines = [
            f"┌{'─'*68}┐",
            f"│ {'历史记录列表':^66} │",
            f"├{'─'*68}┤",
        ]

        for i, r in enumerate(records, 1):
            status = "✅" if r.percentage >= 60 else "⚠️"
            date_str = r.created_at.strftime('%m-%d %H:%M')
            name = r.filename[:25] + "..." if len(r.filename) > 25 else r.filename
            lines.append(f"│ {i:>3}. {status} {name:<30} {r.total_score:>3}/{r.total_max:<3} ({r.percentage:>5.1f}%)  {date_str} │")

        lines.append(f"└{'─'*68}┘")
        return "\n".join(lines)

    def export_history_to_excel(self, user_id: int, output_path: Path, sort_order: str = "desc") -> bool:
        """导出用户所有历史记录为Excel"""
        records = self.db.get_user_records(user_id, limit=10000, offset=0)
        # 根据排序参数调整顺序
        if sort_order == "asc":
            records = list(reversed(records))
        if not records:
            logger.warning("没有历史记录可导出")
            return False

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "历史记录"

        # 标题样式
        title_font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
        title_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        title_align = Alignment(horizontal="center", vertical="center")

        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        cell_font = Font(name="微软雅黑", size=10)
        cell_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell_align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # 标题
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = "实验报告批阅历史记录"
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = title_align
        ws.row_dimensions[1].height = 35

        # 日期
        ws.merge_cells("A2:K2")
        date_cell = ws["A2"]
        date_cell.value = f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | 用户ID: {user_id}"
        date_cell.font = Font(name="微软雅黑", size=10, italic=True, color="666666")
        date_cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[2].height = 25

        # 表头
        headers = ["ID", "文件名", "总分", "满分", "得分率", "风格", "质量分", "实验目的", "实验步骤", "结果分析", "结论", "评语摘要", "建议摘要", "优点", "不足", "创建时间"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        ws.row_dimensions[4].height = 30

        # 数据行
        for i, r in enumerate(records, 1):
            row = i + 4
            dims = r.dimensions

            dim_cols = {
                "实验目的理解": "",
                "实验步骤完整性": "",
                "结果分析与数据处理": "",
                "结论与总结": ""
            }
            for name, dim in dims.items():
                if name in dim_cols:
                    dim_cols[name] = f"{dim.get('score', 0)}/{dim.get('max_score', 0)}"

            values = [
                r.display_id,
                r.filename,
                r.total_score,
                r.total_max,
                f"{r.percentage}%",
                r.style_name,
                f"{r.quality_score:.1f}" if r.quality_score else "-",
                dim_cols["实验目的理解"],
                dim_cols["实验步骤完整性"],
                dim_cols["结果分析与数据处理"],
                dim_cols["结论与总结"],
                r.overall_comment[:200] if r.overall_comment else "",
                r.overall_suggestions[:200] if r.overall_suggestions else "",
                ", ".join(r.strengths[:5]) if r.strengths else "",
                ", ".join(r.weaknesses[:5]) if r.weaknesses else "",
                r.created_at.strftime('%Y-%m-%d %H:%M:%S')
            ]

            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=val)
                cell.font = cell_font
                cell.alignment = cell_align_center if col in [1, 3, 4, 5, 7, 8, 9, 10, 11] else cell_align
                cell.border = thin_border

            ws.row_dimensions[row].height = 60

        # 列宽
        col_widths = [6, 25, 8, 8, 10, 12, 10, 12, 12, 14, 12, 35, 35, 30, 30, 18]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[chr(64 + i) if i <= 26 else 'A' + chr(64 + i - 26)].width = w

        # 统计sheet
        ws_stats = wb.create_sheet("统计汇总")
        stats = self.db.get_user_statistics(user_id)
        stats_data = [
            ["统计项", "数值"],
            ["总记录数", stats['total_count']],
            ["平均分", f"{stats['avg_percentage']:.1f}%"],
            ["最高分", f"{stats['max_percentage']:.1f}%"],
            ["最低分", f"{stats['min_percentage']:.1f}%"],
            ["及格数", stats['pass_count']],
            ["优秀数", stats['excellent_count']],
        ]
        for row_idx, row_data in enumerate(stats_data, 1):
            for col_idx, val in enumerate(row_data, 1):
                cell = ws_stats.cell(row=row_idx, column=col_idx, value=val)
                cell.font = header_font if row_idx == 1 else cell_font
                cell.fill = header_fill if row_idx == 1 else PatternFill()
                cell.alignment = cell_align_center
                cell.border = thin_border

        ws_stats.column_dimensions['A'].width = 15
        ws_stats.column_dimensions['B'].width = 15

        wb.save(output_path)
        logger.info(f"历史记录已导出: {output_path}")
        return True

    def get_record_for_regrading(self, user_id: int, filename: str) -> Optional[GradingRecord]:
        """获取用于重新批阅的记录"""
        return self.db.get_record_by_filename(user_id, filename)
    
    def delete_records(self, record_ids: list[int]) -> dict:
        """批量删除记录"""
        results = {"success": [], "failed": []}
        for rid in record_ids:
            if self.db.delete_record(rid):
                results["success"].append(rid)
            else:
                results["failed"].append(rid)
        return results

    def create_manual_record(self, user_id: int, data: dict) -> int:
        """手动新增一条记录"""
        score_data = {
            "total_score": data.get("total_score", 0),
            "total_max": data.get("total_max", 100),
            "percentage": data.get("percentage", 0),
            "dimensions": data.get("dimensions", {}),
            "overall_comment": data.get("overall_comment", ""),
            "overall_suggestions": data.get("overall_suggestions", ""),
            "strengths": data.get("strengths", []),
            "weaknesses": data.get("weaknesses", []),
        }
        
        return self.db.save_grading_record(
            user_id=user_id,
            filename=data.get("filename", "手动添加"),
            file_path=data.get("file_path", ""),
            score_data=score_data,
            ocr_text=data.get("ocr_text", ""),
            sections=data.get("sections", {}),
            quality_score=data.get("quality_score"),
            style_key=data.get("style_key", "standard"),
            style_name=data.get("style_name", "标准风格"),
            is_aborted=False
        )